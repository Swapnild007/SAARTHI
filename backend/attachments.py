from __future__ import annotations

import base64
import csv
import io
import json
import os
from typing import Any

MAX_FILE_BYTES = int(os.getenv("SAARTHI_MAX_ATTACHMENT_BYTES", str(3 * 1024 * 1024)))
MAX_TEXT_CHARS = int(os.getenv("SAARTHI_MAX_ATTACHMENT_TEXT", "50000"))
IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif"}

def _b64_bytes(value: str) -> bytes:
    raw = value.split(",", 1)[1] if "," in value and value.startswith("data:") else value
    return base64.b64decode(raw)

def inspect_attachment(item: dict[str, Any]) -> dict[str, Any]:
    name = str(item.get("name") or "attachment")
    mime = str(item.get("type") or "application/octet-stream").lower()
    data = str(item.get("data") or "")
    if not data:
        return {"name": name, "type": mime, "error": "empty attachment"}
    try:
        blob = _b64_bytes(data)
    except Exception:
        return {"name": name, "type": mime, "error": "invalid base64 attachment"}
    if len(blob) > MAX_FILE_BYTES:
        return {"name": name, "type": mime, "error": f"attachment exceeds {MAX_FILE_BYTES // 1024 // 1024} MB limit"}
    result: dict[str, Any] = {
        "name": name,
        "type": mime,
        "size": len(blob),
        "kind": "image" if mime in IMAGE_TYPES else "file",
    }
    if mime in IMAGE_TYPES:
        # Keep the original data URL available to a vision-capable provider.
        result["data_url"] = data if data.startswith("data:") else f"data:{mime};base64,{data}"
        result["summary"] = f"Image attachment: {name} ({len(blob):,} bytes)."
        return result

    suffix = os.path.splitext(name.lower())[1]
    text: str | None = None
    if mime in {"text/plain", "text/csv", "text/tab-separated-values", "application/json"} or suffix in {".txt", ".md", ".csv", ".tsv", ".json"}:
        try:
            text = blob.decode("utf-8", errors="replace")
        except Exception:
            text = None
    elif suffix in {".xlsx", ".xlsm"} or mime in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel.sheet.macroenabled.12"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            sheets = {}
            for ws in wb.worksheets[:10]:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row[:30]]
                    if any(vals):
                        rows.append(vals)
                    if len(rows) >= 100:
                        break
                sheets[ws.title] = rows
            result["workbook"] = {"sheets": list(wb.sheetnames), "preview": sheets}
            result["summary"] = f"Excel workbook with {len(wb.sheetnames)} sheet(s): {', '.join(wb.sheetnames[:10])}."
            return result
        except Exception as exc:
            result["error"] = f"Excel parsing unavailable: {exc}"
            return result
    elif suffix == ".pdf" or mime == "application/pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(blob))
            pages = []
            for page in reader.pages[:20]:
                pages.append(page.extract_text() or "")
            text = "\n\n".join(pages)
            result["pages"] = len(reader.pages)
        except Exception as exc:
            result["error"] = f"PDF text extraction unavailable: {exc}"
            return result

    if text is not None:
        text = text[:MAX_TEXT_CHARS]
        result["text"] = text
        if suffix in {".csv", ".tsv"} or mime in {"text/csv", "text/tab-separated-values"}:
            delimiter = "\t" if suffix == ".tsv" or mime == "text/tab-separated-values" else ","
            try:
                rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
                result["tabular"] = {
                    "columns": rows[0][:50] if rows else [],
                    "rows": len(rows),
                    "preview": rows[1:21],
                }
            except Exception:
                pass
        elif suffix == ".json" or mime == "application/json":
            try:
                parsed = json.loads(text)
                result["json_shape"] = type(parsed).__name__
                if isinstance(parsed, list):
                    result["rows"] = len(parsed)
            except Exception:
                pass
        result["summary"] = f"Text/document attachment: {name}, {len(text):,} characters extracted."
    else:
        result["summary"] = f"Binary attachment: {name} ({len(blob):,} bytes)."
    return result

def build_tabular_dataset(item: dict[str, Any]) -> dict[str, Any]:
    """Return a bounded, JSON-safe dataset for deterministic analysis."""
    tab = item.get("tabular") if isinstance(item, dict) else None
    if isinstance(tab, dict) and isinstance(tab.get("columns"), list) and isinstance(tab.get("preview"), list):
        return {
            "name": item.get("name", "dataset"),
            "columns": [str(c) for c in tab["columns"][:50]],
            "rows": tab["preview"][:100],
            "row_count": len(tab["preview"]),
            "source": "attachment-preview",
        }
    return {"name": item.get("name", "dataset"), "columns": [], "rows": [], "row_count": 0, "source": "unsupported"}


def inspect_attachments(items: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    return [inspect_attachment(item) for item in (items or [])][:10]

def provider_content_parts(message: str, inspected: list[dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    text_sections = [message]
    parts: list[dict[str, Any]] = [{"type": "text", "text": message}]
    for item in inspected:
        if item.get("error"):
            text_sections.append(f"[Attachment {item.get('name')}: {item['error']}]")
            continue
        if item.get("kind") == "image" and item.get("data_url"):
            parts.append({"type": "image_url", "image_url": {"url": item["data_url"]}})
            text_sections.append(f"[Image attached: {item.get('name')}]")
            continue
        text_sections.append(f"[Attachment: {item.get('name')}]\n{item.get('summary','')}")
        if item.get("text"):
            text_sections.append(item["text"][:MAX_TEXT_CHARS])
        if item.get("tabular"):
            text_sections.append(json.dumps(item["tabular"], ensure_ascii=False))
        if item.get("workbook"):
            text_sections.append(json.dumps(item["workbook"], ensure_ascii=False))
        if item.get("json_shape"):
            text_sections.append(f"JSON shape: {item['json_shape']}; rows: {item.get('rows','unknown')}")
    combined = "\n\n".join(text_sections)
    if len(combined) > MAX_TEXT_CHARS * 2:
        combined = combined[:MAX_TEXT_CHARS * 2] + "\n[Attachment context truncated]"
    parts[0]["text"] = combined
    return combined, parts
